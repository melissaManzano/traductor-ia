"""
Extraccion de texto de documentos (PDF, Word, TXT) - req. 1.4.

Se aplica herencia y polimorfismo: `DocumentParser` define el contrato
`extract_text()` y cada subclase sabe leer su propio formato. El endpoint
no necesita saber que tipo de archivo llego; solo pide el parser correcto
a la fabrica `get_parser_for_extension`.
"""

from abc import ABC, abstractmethod
import io

from .exceptions import EmptyDocumentError


class DocumentParser(ABC):
    """Clase base abstracta para todo parser de documentos."""

    @abstractmethod
    def extract_text(self, file_bytes: bytes) -> str:
        """Devuelve el texto plano extraido del documento."""
        raise NotImplementedError

    def extract_and_validate(self, file_bytes: bytes) -> str:
        text = self.extract_text(file_bytes)
        if not text or not text.strip():
            raise EmptyDocumentError(
                "El documento no contiene texto procesable (puede estar vacio, "
                "ser una imagen escaneada sin OCR, o estar danado)."
            )
        return text


class PDFParser(DocumentParser):
    def extract_text(self, file_bytes: bytes) -> str:
        from pypdf import PdfReader  # import perezoso: solo se carga si se usa

        reader = PdfReader(io.BytesIO(file_bytes))
        parts = []
        for page in reader.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                parts.append(page_text.strip())
        return "\n\n".join(parts)


class WordParser(DocumentParser):
    def extract_text(self, file_bytes: bytes) -> str:
        import docx  # python-docx

        document = docx.Document(io.BytesIO(file_bytes))
        parts = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text.strip())
        # Tambien se recorre el contenido de las tablas, si el documento tiene.
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)
        return "\n\n".join(parts)


class TextParser(DocumentParser):
    def extract_text(self, file_bytes: bytes) -> str:
        return file_bytes.decode("utf-8", errors="ignore")


_PARSERS: dict[str, type[DocumentParser]] = {
    "pdf": PDFParser,
    "docx": WordParser,
    "txt": TextParser,
}


def get_parser_for_extension(extension: str) -> DocumentParser:
    parser_cls = _PARSERS.get(extension)
    if parser_cls is None:
        # Este caso ya deberia haber sido detectado por FileValidator, pero
        # se protege igual para evitar que la fabrica reciba algo inesperado.
        raise ValueError(f"No existe parser para la extension: {extension}")
    return parser_cls()
